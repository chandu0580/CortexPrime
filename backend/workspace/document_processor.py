"""
Document Processor — extracts raw text from uploaded files.

Supported:
  PDF   — pymupdf (fitz) with page-level granularity
  DOCX  — python-docx
  PPTX  — python-pptx
  TXT   — plain read
  MD    — plain read (markdown source preserved)
  Image — pytesseract OCR (Pillow)
"""
from __future__ import annotations

import hashlib
import io
import logging
from typing import Dict, List, Tuple

logger = logging.getLogger(__name__)


# =========================================================
# PAGE-LEVEL RESULT
# =========================================================

class PageContent:
    """Text + metadata for a single logical page / slide / section."""

    __slots__ = ("page_number", "text", "metadata")

    def __init__(self, page_number: int, text: str, metadata: Dict = None):
        self.page_number = page_number
        self.text        = text.strip()
        self.metadata    = metadata or {}


# =========================================================
# EXTRACTOR
# =========================================================

class DocumentProcessor:
    """Extract text from uploaded files, returning a list of PageContent."""

    # ── Public API ──────────────────────────────────────────────────

    def process_bytes(
        self, data: bytes, filename: str
    ) -> Tuple[List[PageContent], str]:
        """
        Returns (pages, content_hash).

        ``pages``        — ordered list of PageContent
        ``content_hash`` — SHA-256 hex of raw bytes
        """
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else "txt"
        content_hash = hashlib.sha256(data).hexdigest()

        try:
            if ext == "pdf":
                pages = self._pdf(data)
            elif ext == "docx":
                pages = self._docx(data)
            elif ext == "pptx":
                pages = self._pptx(data)
            elif ext in ("txt", "md"):
                pages = self._plaintext(data)
            elif ext in ("png", "jpg", "jpeg", "gif", "webp", "bmp", "tiff"):
                pages = self._image(data)
            else:
                # Attempt as plain text
                pages = self._plaintext(data)
        except Exception as exc:
            logger.error(f"Document processing failed for {filename}: {exc}")
            pages = [PageContent(1, f"[Extraction error: {exc}]")]

        return pages, content_hash

    # ── PDF (pymupdf) ───────────────────────────────────────────────

    def _pdf(self, data: bytes) -> List[PageContent]:
        import fitz  # pymupdf
        doc    = fitz.open(stream=data, filetype="pdf")
        pages  = []
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text")
            if text.strip():
                pages.append(PageContent(
                    page_number = i,
                    text        = text,
                    metadata    = {
                        "width":  page.rect.width,
                        "height": page.rect.height,
                    },
                ))
        doc.close()
        return pages or [PageContent(1, "[PDF contained no extractable text]")]

    # ── DOCX (python-docx) ──────────────────────────────────────────

    def _docx(self, data: bytes) -> List[PageContent]:
        from docx import Document
        doc        = Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Approximate page boundaries: ~40 paragraphs per page
        pages: List[PageContent] = []
        chunk_size = 40
        for i, start in enumerate(range(0, len(paragraphs), chunk_size), start=1):
            text = "\n".join(paragraphs[start : start + chunk_size])
            pages.append(PageContent(i, text))
        return pages or [PageContent(1, "[DOCX contained no extractable text]")]

    # ── PPTX (python-pptx) ─────────────────────────────────────────

    def _pptx(self, data: bytes) -> List[PageContent]:
        from pptx import Presentation
        prs   = Presentation(io.BytesIO(data))
        pages = []
        for i, slide in enumerate(prs.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t:
                            texts.append(t)
            if texts:
                pages.append(PageContent(i, "\n".join(texts), {"slide": i}))
        return pages or [PageContent(1, "[PPTX contained no extractable text]")]

    # ── Plain text / Markdown ───────────────────────────────────────

    def _plaintext(self, data: bytes) -> List[PageContent]:
        for enc in ("utf-8", "latin-1", "cp1252"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = data.decode("utf-8", errors="replace")

        # Split into logical pages (~3000 chars each)
        pages    = []
        page_sz  = 3000
        for i, start in enumerate(range(0, len(text), page_sz), start=1):
            segment = text[start : start + page_sz].strip()
            if segment:
                pages.append(PageContent(i, segment))
        return pages or [PageContent(1, text or "[Empty document]")]

    # ── Image OCR (pytesseract / Pillow) ───────────────────────────

    def _image(self, data: bytes) -> List[PageContent]:
        try:
            from PIL import Image
            import pytesseract
            img  = Image.open(io.BytesIO(data))
            text = pytesseract.image_to_string(img)
            return [PageContent(1, text.strip() or "[No text detected in image]")]
        except ImportError:
            logger.warning("pytesseract not installed — returning placeholder for image")
            return [PageContent(1, "[Image content — OCR unavailable]")]
        except Exception as exc:
            logger.error(f"Image OCR failed: {exc}")
            return [PageContent(1, f"[Image OCR error: {exc}]")]


# =========================================================
# SINGLETON
# =========================================================

document_processor = DocumentProcessor()
