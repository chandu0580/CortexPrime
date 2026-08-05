"""Generate submission-quality Word documents from the markdown sources.

Produces documents that read as professional deliverables rather than printed
markdown:

  * designed cover page with a colour banner and a document-control table
  * auto-updating Table of Contents with page numbers
  * heading hierarchy with accent rules under major sections
  * banded tables with captions ("Table 3: ...")
  * callout boxes with a coloured left accent bar
  * running header with a rule, footer with "Page X of Y"

Edit the ``.md`` source, re-run this, and the ``.docx`` matches. One source
means the two cannot drift apart.

    python submission/make_docx.py
"""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SUBMISSION_DIR = Path(__file__).resolve().parent

# ----------------------------------------------------------------------
# Submission details
# ----------------------------------------------------------------------

PROJECT_NAME = "CortexPrime"
PROJECT_SUBTITLE = "AI-Native Enterprise Operations Intelligence Platform"
REGISTRATION_NO = "24BMMCA023"
STUDENT_NAME = "Chandu S"
DOCUMENT_VERSION = "1.0"

# ----------------------------------------------------------------------
# Design tokens
# ----------------------------------------------------------------------

INK = RGBColor(0x11, 0x18, 0x27)          # body text
ACCENT = RGBColor(0x1E, 0x40, 0xAF)       # headings, rules
MUTED = RGBColor(0x6B, 0x72, 0x80)        # secondary text
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEX_ACCENT = "1E40AF"
HEX_ACCENT_SOFT = "EEF2FF"
HEX_BAND = "F8FAFC"
HEX_CODE = "F1F5F9"
HEX_NOTE = "FFFBEB"
HEX_NOTE_BAR = "F59E0B"
HEX_RULE = "CBD5E1"

BODY_FONT = "Calibri"
MONO_FONT = "Consolas"

DOCUMENTS = [
    (
        "3-Project-Installation-Procedure.md",
        "3-Project-Installation-Procedure.docx",
        "Project Installation Procedure",
        "Document 3",
    ),
    (
        "3.1-Supporting-Software.md",
        "3.1-Supporting-Software.docx",
        "Supporting Software",
        "Document 3.1",
    ),
    (
        "3.3-README.md",
        "3.3-README.docx",
        "Readme — How to Run the Project",
        "Document 3.3",
    ),
]

_BOLD = re.compile(r"\*\*([^*]+)\*\*")
_ITALIC = re.compile(r"(?<!\*)\*([^*]+)\*(?!\*)")
_CODE = re.compile(r"`([^`]+)`")
_TABLE_DIVIDER = re.compile(r"^\s*\|[\s:\-|]+\|\s*$")


# ======================================================================
# Low-level Word primitives
# ======================================================================


def _element(tag: str, **attributes) -> OxmlElement:
    node = OxmlElement(tag)
    for key, value in attributes.items():
        node.set(qn(f"w:{key}"), value)
    return node


def _shade_paragraph(paragraph, fill: str) -> None:
    paragraph._p.get_or_add_pPr().append(
        _element("w:shd", val="clear", color="auto", fill=fill)
    )


def _shade_cell(cell, fill: str) -> None:
    cell._tc.get_or_add_tcPr().append(
        _element("w:shd", val="clear", color="auto", fill=fill)
    )


def _paragraph_border(paragraph, *, edge: str, size: int, colour: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    borders.append(
        _element(f"w:{edge}", val="single", sz=str(size), space="6", color=colour)
    )


def _cell_width(cell, inches: float) -> None:
    cell.width = Inches(inches)
    cell._tc.get_or_add_tcPr().append(
        _element("w:tcW", w=str(int(inches * 1440)), type="dxa")
    )


def _remove_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_element(f"w:{edge}", val="none", sz="0"))
    properties.append(borders)


def _field(paragraph, instruction: str) -> None:
    """Insert a Word field code (PAGE, NUMPAGES, TOC)."""
    run = paragraph.add_run()
    run._r.append(_element("w:fldChar", fldCharType="begin"))

    instruction_node = OxmlElement("w:instrText")
    instruction_node.set(qn("xml:space"), "preserve")
    instruction_node.text = instruction
    run._r.append(instruction_node)

    run._r.append(_element("w:fldChar", fldCharType="separate"))
    run._r.append(_element("w:fldChar", fldCharType="end"))


def _force_field_update_on_open(document: Document) -> None:
    """Make Word refresh the Table of Contents when the file is opened.

    Without this the TOC shows a placeholder until someone presses F9, which is
    exactly the sort of detail that makes a document look unfinished.
    """
    settings = document.settings.element
    if settings.find(qn("w:updateFields")) is None:
        settings.append(_element("w:updateFields", val="true"))


# ======================================================================
# Document chrome
# ======================================================================


def _configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.18

    spec = {1: (17, 22, 8), 2: (13.5, 16, 6), 3: (11.5, 12, 4), 4: (10.5, 10, 3)}
    for level, (size, before, after) in spec.items():
        style = document.styles[f"Heading {level}"]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = document.styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(3)


def _setup_page(document: Document, doc_label: str) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)
    section.different_first_page_header_footer = True

    header = section.header.paragraphs[0]
    header.text = f"{PROJECT_NAME}   |   {doc_label}"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.name = BODY_FONT
        run.font.color.rgb = MUTED
    _paragraph_border(header, edge="bottom", size=6, colour=HEX_RULE)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prefix = footer.add_run("Page ")
    prefix.font.size = Pt(8)
    prefix.font.color.rgb = MUTED
    _field(footer, " PAGE ")
    middle = footer.add_run(" of ")
    middle.font.size = Pt(8)
    middle.font.color.rgb = MUTED
    _field(footer, " NUMPAGES ")


def _banner(document: Document, text: str, subtext: str) -> None:
    """Full-width colour block at the top of the cover page."""
    table = document.add_table(rows=1, cols=1)
    _remove_table_borders(table)
    cell = table.rows[0].cells[0]
    _cell_width(cell, 6.6)
    _shade_cell(cell, HEX_ACCENT)

    title = cell.paragraphs[0]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(26)
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run(text)
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.color.rgb = WHITE
    run.font.name = BODY_FONT

    sub = cell.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(26)
    run = sub.add_run(subtext)
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(0xDB, 0xE4, 0xFF)
    run.font.name = BODY_FONT


def _detail_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _remove_table_borders(table)

    for index, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        _cell_width(cells[0], 1.9)
        _cell_width(cells[1], 3.3)
        if index % 2 == 0:
            _shade_cell(cells[0], HEX_BAND)
            _shade_cell(cells[1], HEX_BAND)

        left = cells[0].paragraphs[0]
        left.paragraph_format.space_before = Pt(5)
        left.paragraph_format.space_after = Pt(5)
        run = left.add_run(label)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = MUTED

        right = cells[1].paragraphs[0]
        right.paragraph_format.space_before = Pt(5)
        right.paragraph_format.space_after = Pt(5)
        run = right.add_run(value)
        run.font.size = Pt(10.5)
        run.font.color.rgb = INK


def _cover_page(document: Document, title: str, label: str) -> None:
    _banner(document, PROJECT_NAME, PROJECT_SUBTITLE)

    for _ in range(4):
        document.add_paragraph()

    tag = document.add_paragraph()
    tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = tag.add_run(label.upper())
    run.font.size = Pt(9.5)
    run.font.bold = True
    run.font.color.rgb = ACCENT
    run.font.name = BODY_FONT

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(10)
    run = heading.add_run(title)
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = INK
    run.font.name = BODY_FONT

    rule = document.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _paragraph_border(rule, edge="bottom", size=12, colour=HEX_ACCENT)

    for _ in range(5):
        document.add_paragraph()

    _detail_table(
        document,
        [
            ("Registration No.", REGISTRATION_NO),
            ("Name", STUDENT_NAME),
            ("Document Version", DOCUMENT_VERSION),
            ("Date", date.today().strftime("%d %B %Y")),
        ],
    )

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _table_of_contents(document: Document) -> None:
    heading = document.add_heading("Table of Contents", level=1)
    _paragraph_border(heading, edge="bottom", size=8, colour=HEX_RULE)

    paragraph = document.add_paragraph()
    _field(paragraph, r'TOC \o "1-3" \h \z \u')

    hint = document.add_paragraph()
    run = hint.add_run(
        "Page numbers update automatically when this document is opened in Word."
    )
    run.font.size = Pt(8.5)
    run.font.italic = True
    run.font.color.rgb = MUTED

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ======================================================================
# Markdown rendering
# ======================================================================


def _plain(text: str) -> str:
    return _ITALIC.sub(r"\1", _BOLD.sub(r"\1", _CODE.sub(r"\1", text)))


def _add_runs(paragraph, text: str) -> None:
    for token in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = MONO_FONT
            run.font.size = Pt(9.5)
            run.font.color.rgb = ACCENT
        elif token.startswith("*") and token.endswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        else:
            paragraph.add_run(token)


def _code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.right_indent = Inches(0.1)
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(11)
    paragraph.paragraph_format.line_spacing = 1.0
    _shade_paragraph(paragraph, HEX_CODE)
    _paragraph_border(paragraph, edge="left", size=18, colour=HEX_ACCENT)

    run = paragraph.add_run("\n".join(lines))
    run.font.name = MONO_FONT
    run.font.size = Pt(9)
    run.font.color.rgb = INK


def _callout(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(10)
    _shade_paragraph(paragraph, HEX_NOTE)
    _paragraph_border(paragraph, edge="left", size=18, colour=HEX_NOTE_BAR)
    _add_runs(paragraph, text)
    for run in paragraph.runs:
        run.font.size = Pt(10)


def _caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(12)
    run = paragraph.add_run(text)
    run.font.size = Pt(8.5)
    run.font.italic = True
    run.font.color.rgb = MUTED


def _data_table(document: Document, rows: list[str], number: int, title: str) -> None:
    parsed = [
        [cell.strip() for cell in row.strip().strip("|").split("|")]
        for row in rows
        if not _TABLE_DIVIDER.match(row)
    ]
    if not parsed:
        return

    width = max(len(row) for row in parsed)
    table = document.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for index, values in enumerate(parsed):
        cells = table.add_row().cells
        for position in range(width):
            text = _plain(values[position]) if position < len(values) else ""
            paragraph = cells[position].paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run(text)
            run.font.size = Pt(9.5)
            run.font.name = BODY_FONT

            if index == 0:
                run.bold = True
                run.font.color.rgb = WHITE
                _shade_cell(cells[position], HEX_ACCENT)
            elif index % 2 == 0:
                _shade_cell(cells[position], HEX_BAND)

    _caption(document, f"Table {number}: {title}")


def _render(document: Document, source: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()

    in_code = False
    code: list[str] = []
    table: list[str] = []
    table_number = 0
    last_heading = "Reference"
    front_matter = True
    first_body_heading = True

    def flush_table() -> None:
        nonlocal table_number
        if table:
            table_number += 1
            _data_table(document, list(table), table_number, last_heading)
            table.clear()

    for line in lines:
        stripped = line.strip()

        # Cover page already carries the title and submission details.
        if front_matter:
            if stripped.startswith("## ") or (
                stripped.startswith("# ") and not front_matter
            ):
                front_matter = False
            elif (
                not stripped
                or stripped == "---"
                or stripped.startswith("# ")
                or stripped.startswith("**Project:**")
                or stripped.startswith("**Reg. No")
            ):
                continue
            else:
                front_matter = False

        if stripped.startswith("```"):
            flush_table()
            if in_code:
                _code_block(document, code)
                code.clear()
            in_code = not in_code
            continue

        if in_code:
            code.append(line)
            continue

        if stripped.startswith("|"):
            table.append(line)
            continue
        flush_table()

        if not stripped or stripped == "---":
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = _plain(stripped.lstrip("#").strip())
            level = min(level, 4)

            if level == 1 and not first_body_heading:
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            first_body_heading = False

            heading = document.add_heading(text, level=level)
            if level <= 2:
                _paragraph_border(heading, edge="bottom", size=6, colour=HEX_RULE)
            last_heading = text
            continue

        if stripped.startswith(">"):
            _callout(document, stripped.lstrip("> ").strip())
            continue

        if re.match(r"^\d+\.\s", stripped):
            paragraph = document.add_paragraph(style="List Number")
            _add_runs(paragraph, re.sub(r"^\d+\.\s", "", stripped))
            continue

        if stripped.startswith(("- ", "* ")):
            paragraph = document.add_paragraph(style="List Bullet")
            _add_runs(paragraph, stripped[2:])
            continue

        _add_runs(document.add_paragraph(), stripped)

    flush_table()


# ======================================================================
# Entry point
# ======================================================================


def convert(source: Path, destination: Path, title: str, label: str) -> None:
    document = Document()
    _configure_styles(document)
    _setup_page(document, title)
    _cover_page(document, title, label)
    _table_of_contents(document)
    _render(document, source)
    _force_field_update_on_open(document)
    document.save(destination)
    print(f"  {destination.name:<44} {destination.stat().st_size / 1024:>6.1f} KB")


def main() -> None:
    print(f"\nGenerating submission documents for {STUDENT_NAME} ({REGISTRATION_NO})\n")
    for source_name, output_name, title, label in DOCUMENTS:
        source = SUBMISSION_DIR / source_name
        if not source.exists():
            print(f"  skipped {source_name} (not found)")
            continue
        convert(source, SUBMISSION_DIR / output_name, title, label)
    print(
        "\nDone. Open each file in Word — the Table of Contents fills in "
        "automatically on first open."
    )


if __name__ == "__main__":
    main()
