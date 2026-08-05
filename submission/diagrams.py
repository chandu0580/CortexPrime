"""Vector diagram helpers for the project presentation.

Everything here draws with native PowerPoint shapes rather than embedded
images, so a diagram stays crisp at any zoom and can be edited in PowerPoint if
a label needs changing during the viva.

Colours are taken from the running CortexPrime UI — the tokens in
``frontend/app/globals.css`` under the default dark theme — so the deck and the
product look like the same thing.

    cycle           circular process loop
    flowchart       decision flow with a diamond branch
    hash_chain      linked blocks showing tamper-evidence
    layer_stack     architecture layers with a dependency arrow
    split_compare   before / after comparison
    data_flow       sources -> process -> sinks   (not in the current deck)
    use_case        actors and system boundary    (not in the current deck)
"""

from __future__ import annotations

import math

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ----------------------------------------------------------------------
# Palette — CortexPrime dark theme (frontend/app/globals.css, :root)
# ----------------------------------------------------------------------

BG            = RGBColor(0x0A, 0x10, 0x0D)   # --background
SURFACE       = RGBColor(0x12, 0x1A, 0x16)   # --surface
RAISED        = RGBColor(0x1A, 0x25, 0x20)   # --surface-raised
OVERLAY       = RGBColor(0x22, 0x30, 0x28)   # --surface-overlay
BORDER        = RGBColor(0x20, 0x26, 0x23)   # --border   over --background
BORDER_STRONG = RGBColor(0x31, 0x36, 0x34)   # --border-strong

TEXT     = RGBColor(0xF0, 0xF7, 0xF4)        # --text-primary
TEXT_DIM = RGBColor(0xA6, 0xAD, 0xAA)        # --text-secondary
MUTED    = RGBColor(0x6B, 0x71, 0x6E)        # --text-muted

ACCENT    = RGBColor(0x82, 0xC0, 0xA4)       # --accent-primary
ACCENT_HI = RGBColor(0x96, 0xCE, 0xAD)       # --accent-hover
ACCENT_2  = RGBColor(0x4A, 0x8C, 0x70)       # --accent-secondary
SUCCESS   = RGBColor(0x38, 0xB8, 0x8A)       # --success
WARNING   = RGBColor(0xF9, 0xA8, 0x25)       # --warning
DANGER    = RGBColor(0xF8, 0x71, 0x71)       # --danger

FONT = "Segoe UI"
# ✗ ✓ ▸ are absent from Segoe UI itself and only live in Segoe UI Symbol. Naming
# it explicitly avoids relying on PowerPoint's fallback, which substitutes a
# different face and renders the marks at an inconsistent weight.
SYMBOL_FONT = "Segoe UI Symbol"


def ink_for(fill: RGBColor) -> RGBColor:
    """Pick dark or light text for a fill, by perceived luminance.

    The sage palette straddles the boundary — ``#82c0a4`` needs dark text while
    ``#4a8c70`` needs light — so this is decided per colour rather than guessed
    at each call site.
    """
    r, g, b = fill[0], fill[1], fill[2]
    luminance = (0.299 * r + 0.587 * g + 0.114 * b) / 255
    return BG if luminance > 0.55 else TEXT


def _dim_for(fill: RGBColor) -> RGBColor:
    """The secondary-text partner to :func:`ink_for`, for sub-labels.

    On a dark fill this is a desaturated sage rather than ``TEXT_DIM`` — the
    theme's secondary grey is tuned for the near-black page background and goes
    muddy against the mid-green surfaces used inside diagrams.
    """
    return RGBColor(0x1E, 0x33, 0x29) if ink_for(fill) == BG else RGBColor(0xC6, 0xDB, 0xD1)


# ======================================================================
# Primitives
# ======================================================================


def label(shape, lines, *, size=12, colour=TEXT, bold=True, align=PP_ALIGN.CENTER):
    """Write centred text inside a shape."""
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.04)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    # Ovals and diamonds are narrow at the top, so top-anchored text spills
    # outside the drawn outline. Centre it vertically in every shape.
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    for index, item in enumerate(lines):
        text, item_size, item_bold, item_colour = (
            item if isinstance(item, tuple) else (item, size, bold, colour)
        )
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(1)
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(item_size)
        run.font.bold = item_bold
        run.font.color.rgb = item_colour
        run.font.name = FONT


def node(slide, left, top, width, height, *, fill=ACCENT, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         line=None, lines=(), size=12, colour=None):
    element = slide.shapes.add_shape(shape, left, top, width, height)
    element.fill.solid()
    element.fill.fore_color.rgb = fill
    if line is None:
        element.line.fill.background()
    else:
        element.line.color.rgb = line
        element.line.width = Pt(1.25)
    element.shadow.inherit = False
    if lines:
        label(element, lines, size=size, colour=colour or ink_for(fill))
    return element


def arrow(slide, x1, y1, x2, y2, *, colour=ACCENT_2, width=1.75, dashed=False):
    """A straight connector with an arrowhead at the end.

    Coordinates are ``Inches(...)`` values, given as four separate arguments
    rather than two points — every call site reads more clearly that way.
    """
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    connector.line.color.rgb = colour
    connector.line.width = Pt(width)
    if dashed:
        connector.line.dash_style = 4  # MSO_LINE_DASH_STYLE.DASH

    line_element = connector.line._get_or_add_ln()
    tail = line_element.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "h": "med"})
    line_element.append(tail)
    return connector


def caption(slide, text, *, top, size=11.5, colour=MUTED, left=0.75, width=11.8):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.4))
    label(box, [text], size=size, colour=colour, bold=False, align=PP_ALIGN.CENTER)


# ======================================================================
# Diagram: circular process loop
# ======================================================================


def cycle(slide, steps, *, cx=6.67, cy=3.75, radius=1.75, node_w=2.05, node_h=0.82,
          stretch=2.0, centre_text=None):
    """Steps arranged on a circle with arrows between them.

    Used for the core detect -> explain -> plan -> approve -> execute -> verify
    loop. A circle rather than a line because the loop genuinely repeats.
    """
    count = len(steps)
    positions = []

    for index in range(count):
        # Start at the top and go clockwise.
        angle = -math.pi / 2 + (2 * math.pi * index / count)
        # Stretched horizontally: the slide is 16:9, and on a true circle the
        # 2"-wide nodes sit close enough that the connecting arrows shrink to
        # invisible stubs. Vertical radius is what the slide height limits.
        x = cx + radius * math.cos(angle) * stretch
        y = cy + radius * math.sin(angle)
        positions.append((x, y))

    if centre_text:
        node(
            slide,
            Inches(cx - 1.1), Inches(cy - 0.45), Inches(2.2), Inches(0.9),
            fill=RAISED, line=ACCENT, shape=MSO_SHAPE.OVAL,
            lines=centre_text, size=12, colour=ACCENT,
        )

    # Arrows first so the nodes sit on top of the line ends.
    for index in range(count):
        x1, y1 = positions[index]
        x2, y2 = positions[(index + 1) % count]

        dx, dy = x2 - x1, y2 - y1
        distance = math.hypot(dx, dy) or 1

        # Where the centre-to-centre line leaves the box, as a fraction of its
        # length. Insetting by half the box WIDTH regardless of direction (the
        # obvious shortcut) overshoots badly on the near-vertical legs, where
        # the line exits through the much closer top or bottom edge instead.
        span_x = (node_w / 2) / abs(dx) if dx else float("inf")
        span_y = (node_h / 2) / abs(dy) if dy else float("inf")
        edge = min(span_x, span_y)

        start = edge + 0.10 / distance
        end = 1 - (edge + 0.16 / distance)

        arrow(
            slide,
            Inches(x1 + dx * start), Inches(y1 + dy * start),
            Inches(x1 + dx * end), Inches(y1 + dy * end),
            colour=ACCENT, width=2,
        )

    for index, (title, detail) in enumerate(steps):
        x, y = positions[index]
        fill = ACCENT if index % 2 == 0 else ACCENT_2
        node(
            slide,
            Inches(x - node_w / 2), Inches(y - node_h / 2), Inches(node_w), Inches(node_h),
            fill=fill,
            lines=[(title, 12.5, True, ink_for(fill)), (detail, 9, False, _dim_for(fill))],
        )


# ======================================================================
# Diagram: decision flowchart
# ======================================================================


def flowchart(slide, *, top=1.85):
    """Detection -> risk decision -> approval or auto-execute -> verify."""
    box_w, box_h = 2.15, 0.72
    column = 6.67   # centre column — the middle of a 13.333" slide

    def centre(x, y, w=box_w, h=box_h):
        return Inches(x - w / 2), Inches(y - h / 2), Inches(w), Inches(h)

    y = top + 0.36
    node(slide, *centre(column, y), fill=SUCCESS, shape=MSO_SHAPE.OVAL,
         lines=[("PROBLEM DETECTED", 11, True, ink_for(SUCCESS))])

    y2 = y + 0.95
    arrow(slide, Inches(column), Inches(y + box_h / 2), Inches(column), Inches(y2 - box_h / 2),
          colour=ACCENT)
    node(slide, *centre(column, y2), fill=ACCENT_2,
         lines=[("Gather evidence", 11.5, True, TEXT), ("logs · metrics · history", 8.5, False, TEXT_DIM)])

    y3 = y2 + 0.95
    arrow(slide, Inches(column), Inches(y2 + box_h / 2), Inches(column), Inches(y3 - box_h / 2),
          colour=ACCENT)
    node(slide, *centre(column, y3), fill=ACCENT_2,
         lines=[("Identify root cause", 11.5, True, TEXT), ("with citations", 8.5, False, TEXT_DIM)])

    # Decision diamond
    y4 = y3 + 1.15
    arrow(slide, Inches(column), Inches(y3 + box_h / 2), Inches(column), Inches(y4 - 0.52),
          colour=ACCENT)
    node(slide, Inches(column - 1.25), Inches(y4 - 0.52), Inches(2.5), Inches(1.04),
         fill=WARNING, shape=MSO_SHAPE.DIAMOND,
         lines=[("Risk above LOW?", 10.5, True, ink_for(WARNING))])

    # Right branch — approval required
    right = column + 3.5
    arrow(slide, Inches(column + 1.25), Inches(y4), Inches(right - box_w / 2), Inches(y4),
          colour=DANGER)
    caption_box = slide.shapes.add_textbox(
        Inches(column + 1.35), Inches(y4 - 0.42), Inches(1.4), Inches(0.3))
    label(caption_box, ["YES"], size=10, colour=DANGER, align=PP_ALIGN.LEFT)
    node(slide, *centre(right, y4), fill=DANGER,
         lines=[("Wait for human", 11, True, ink_for(DANGER)),
                ("approval required", 8.5, False, ink_for(DANGER))])

    # Left branch — auto approve
    left = column - 3.5
    arrow(slide, Inches(column - 1.25), Inches(y4), Inches(left + box_w / 2), Inches(y4),
          colour=SUCCESS)
    caption_box = slide.shapes.add_textbox(
        Inches(column - 2.6), Inches(y4 - 0.42), Inches(1.2), Inches(0.3))
    label(caption_box, ["NO"], size=10, colour=SUCCESS, align=PP_ALIGN.RIGHT)
    node(slide, *centre(left, y4), fill=SUCCESS,
         lines=[("Auto-approve", 11, True, ink_for(SUCCESS)),
                ("low risk only", 8.5, False, ink_for(SUCCESS))])

    # Converge on execute
    y5 = y4 + 1.15
    node(slide, *centre(column, y5), fill=ACCENT,
         lines=[("EXECUTE THE FIX", 11.5, True, ink_for(ACCENT)),
                ("digest verified first", 8.5, False, _dim_for(ACCENT))])
    arrow(slide, Inches(right), Inches(y4 + 0.52), Inches(column + box_w / 2), Inches(y5),
          colour=DANGER)
    arrow(slide, Inches(left), Inches(y4 + 0.52), Inches(column - box_w / 2), Inches(y5),
          colour=SUCCESS)

    # Verify
    y6 = y5 + 0.85
    arrow(slide, Inches(column), Inches(y5 + box_h / 2), Inches(column), Inches(y6 - 0.3),
          colour=ACCENT)
    node(slide, Inches(column - 1.4), Inches(y6 - 0.3), Inches(2.8), Inches(0.62),
         fill=SUCCESS, shape=MSO_SHAPE.OVAL,
         lines=[("VERIFY & AUDIT", 11, True, ink_for(SUCCESS))])


# ======================================================================
# Diagram: hash chain
# ======================================================================


def hash_chain(slide, *, top=2.35, compact=False):
    """Linked audit blocks, with one shown as tampered and detected.

    ``compact`` drops the trailing explanation paragraph and shrinks the blocks,
    for when the diagram shares a slide with something else.
    """
    blocks = [
        ("Record 1", "seq 0", "hash a1f3…", False),
        ("Record 2", "seq 1", "prev a1f3… → 7c92…", False),
        ("Record 3", "seq 2", "prev 7c92… → e4b1…", True),
        ("Record 4", "seq 3", "prev e4b1… → 9d20…", False),
    ]
    width, gap = 2.55, 0.55
    height = 1.25 if compact else 1.35
    start = 0.95

    for index, (name, sequence, digest, tampered) in enumerate(blocks):
        x = start + index * (width + gap)
        fill = DANGER if tampered else ACCENT_2
        node(slide, Inches(x), Inches(top), Inches(width), Inches(height),
             fill=fill,
             lines=[
                 (name, 12.5, True, ink_for(fill)),
                 (sequence, 9, False, _dim_for(fill)),
                 (digest, 9, False, _dim_for(fill)),
             ])
        if index < len(blocks) - 1:
            arrow(slide, Inches(x + width), Inches(top + height / 2),
                  Inches(x + width + gap), Inches(top + height / 2),
                  colour=ACCENT, width=2)

    # Callout directly beneath the tampered block. No pointer line: one would
    # have to run straight through this text to reach the block above it, and
    # the red fill already draws the eye without help.
    x_bad = start + 2 * (width + gap)
    box = slide.shapes.add_textbox(Inches(x_bad + width / 2 - 1.7),
                                   Inches(top + height + 0.1), Inches(3.4), Inches(0.6))
    label(box, [
        ("EDITED AFTER THE FACT", 10.5, True, DANGER),
        ("its digest no longer matches", 9.5, False, MUTED),
    ], align=PP_ALIGN.CENTER)

    if compact:
        return

    box = slide.shapes.add_textbox(Inches(0.95), Inches(top + height + 1.05),
                                   Inches(11.5), Inches(0.9))
    label(box, [
        ("Every later record breaks too", 13.5, True, TEXT),
        ("Each digest covers the one before it, so a single edit invalidates the whole "
         "remaining chain. Deleting a record leaves a gap in the sequence numbers — "
         "which a simple link check would miss.", 11.5, False, TEXT_DIM),
    ], align=PP_ALIGN.CENTER)


# ======================================================================
# Diagram: layered architecture with dependency arrow
# ======================================================================


def layer_stack(slide, layers, *, top=1.8, height=0.72, gap=0.12):
    for index, (name, detail, fill) in enumerate(layers):
        y = top + index * (height + gap)
        node(slide, Inches(1.55), Inches(y), Inches(9.4), Inches(height),
             fill=fill,
             lines=[(name, 13, True, ink_for(fill)), (detail, 10, False, _dim_for(fill))])

    total = len(layers) * (height + gap) - gap
    arrow(slide, Inches(11.35), Inches(top + 0.15), Inches(11.35), Inches(top + total - 0.1),
          colour=ACCENT, width=2.25)
    box = slide.shapes.add_textbox(Inches(11.5), Inches(top + total / 2 - 0.45),
                                   Inches(1.7), Inches(0.9))
    label(box, [
        ("depends", 10.5, True, ACCENT),
        ("downward", 10.5, True, ACCENT),
        ("only", 10.5, True, ACCENT),
    ], align=PP_ALIGN.LEFT)


# ======================================================================
# Diagram: before / after comparison
# ======================================================================


def split_compare(slide, left_title, left_items, right_title, right_items, *, top=1.9,
                  panel_h=4.2):
    panel_w = 5.7

    for offset, (title, items, colour) in enumerate(
        [(left_title, left_items, DANGER), (right_title, right_items, SUCCESS)]
    ):
        x = 0.75 + offset * (panel_w + 0.4)
        node(slide, Inches(x), Inches(top), Inches(panel_w), Inches(panel_h),
             fill=SURFACE, line=BORDER_STRONG)
        node(slide, Inches(x), Inches(top), Inches(panel_w), Inches(0.62),
             fill=colour, lines=[(title, 13.5, True, ink_for(colour))])

        box = slide.shapes.add_textbox(
            Inches(x + 0.3), Inches(top + 0.85), Inches(panel_w - 0.55), Inches(panel_h - 1.0))
        frame = box.text_frame
        frame.word_wrap = True
        # Spread the items down the panel instead of stacking them at the top.
        # There are len(items) lines but only len(items) - 1 gaps between them,
        # and the space below the last line is dead — so divide by the gaps.
        line_height = 14 * 1.22 / 72
        free = (panel_h - 1.05) - len(items) * line_height
        spacing = max(free / max(len(items) - 1, 1), 0.08) * 72
        for index, item in enumerate(items):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.space_after = Pt(spacing)
            mark = paragraph.add_run()
            mark.text = f"{'✗' if offset == 0 else '✓'}   "
            mark.font.size = Pt(13)
            mark.font.bold = True
            mark.font.color.rgb = colour
            mark.font.name = SYMBOL_FONT
            run = paragraph.add_run()
            run.text = item
            run.font.size = Pt(14)
            run.font.color.rgb = TEXT
            run.font.name = FONT


# ======================================================================
# Diagram: data flow  (kept available; not used by the current deck)
# ======================================================================


def data_flow(slide, *, top=1.95):
    """External sources -> CortexPrime core -> outputs."""
    sources = [
        ("Docker", "containers, events"),
        ("GitHub / GitLab", "commits, alerts"),
        ("Prometheus", "metrics"),
        ("PostgreSQL", "cost records"),
    ]
    outputs = [
        ("Jira", "incident tickets"),
        ("Dashboard", "live status"),
        ("Audit Trail", "tamper-evident"),
        ("Fix Actions", "applied changes"),
    ]

    col_w, col_h, gap = 2.3, 0.68, 0.22

    box = slide.shapes.add_textbox(Inches(0.75), Inches(top - 0.38), Inches(2.3), Inches(0.3))
    label(box, ["DATA SOURCES"], size=10, colour=MUTED, align=PP_ALIGN.CENTER)
    for index, (name, detail) in enumerate(sources):
        y = top + index * (col_h + gap)
        node(slide, Inches(0.75), Inches(y), Inches(col_w), Inches(col_h),
             fill=SURFACE, line=BORDER_STRONG,
             lines=[(name, 11, True, TEXT), (detail, 8.5, False, MUTED)])
        arrow(slide, Inches(0.75 + col_w), Inches(y + col_h / 2),
              Inches(4.55), Inches(top + 1.55), colour=ACCENT_2, width=1.25)

    node(slide, Inches(4.6), Inches(top + 0.15), Inches(4.1), Inches(2.85),
         fill=ACCENT_2, lines=[])
    box = slide.shapes.add_textbox(Inches(4.75), Inches(top + 0.3), Inches(3.8), Inches(2.6))
    label(box, [
        ("CORTEXPRIME CORE", 13, True, TEXT),
        ("", 6, False, TEXT),
        ("Detect  →  Correlate evidence", 10.5, False, TEXT_DIM),
        ("Identify root cause", 10.5, False, TEXT_DIM),
        ("Classify risk  →  Gate approval", 10.5, False, TEXT_DIM),
        ("Execute  →  Verify outcome", 10.5, False, TEXT_DIM),
    ], align=PP_ALIGN.CENTER)

    box = slide.shapes.add_textbox(Inches(10.25), Inches(top - 0.38), Inches(2.3), Inches(0.3))
    label(box, ["OUTPUTS"], size=10, colour=MUTED, align=PP_ALIGN.CENTER)
    for index, (name, detail) in enumerate(outputs):
        y = top + index * (col_h + gap)
        node(slide, Inches(10.25), Inches(y), Inches(col_w), Inches(col_h),
             fill=SURFACE, line=ACCENT,
             lines=[(name, 11, True, TEXT), (detail, 8.5, False, MUTED)])
        arrow(slide, Inches(8.75), Inches(top + 1.55),
              Inches(10.25), Inches(y + col_h / 2), colour=ACCENT, width=1.25)

    y_human = top + 3.25
    node(slide, Inches(5.55), Inches(y_human), Inches(2.2), Inches(0.6),
         fill=WARNING, shape=MSO_SHAPE.OVAL,
         lines=[("HUMAN APPROVER", 10, True, ink_for(WARNING))])
    arrow(slide, Inches(6.65), Inches(top + 3.0), Inches(6.65), Inches(y_human),
          colour=WARNING, width=1.75)
    arrow(slide, Inches(6.0), Inches(y_human), Inches(6.0), Inches(top + 3.0),
          colour=WARNING, width=1.75)


# ======================================================================
# Diagram: use case  (kept available; not used by the current deck)
# ======================================================================


def use_case(slide, *, top=1.9):
    """Actors outside a system boundary, use cases inside."""
    boundary = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.5), Inches(top), Inches(6.35), Inches(4.35)
    )
    boundary.fill.solid()
    boundary.fill.fore_color.rgb = SURFACE
    boundary.line.color.rgb = ACCENT
    boundary.line.width = Pt(1.5)
    boundary.shadow.inherit = False

    box = slide.shapes.add_textbox(Inches(3.6), Inches(top + 0.1), Inches(6.1), Inches(0.3))
    label(box, ["CORTEXPRIME"], size=11, colour=ACCENT, align=PP_ALIGN.CENTER)

    cases = [
        "Monitor infrastructure",
        "Diagnose root cause",
        "Raise incident ticket",
        "Request approval",
        "Execute approved fix",
        "Verify and audit",
    ]
    for index, text in enumerate(cases):
        row, column = divmod(index, 2)
        node(
            slide,
            Inches(3.75 + column * 3.0), Inches(top + 0.55 + row * 1.22),
            Inches(2.85), Inches(0.92),
            fill=RAISED, line=ACCENT_2, shape=MSO_SHAPE.OVAL,
            lines=[(text, 10.5, False, TEXT)],
        )

    actors = [
        (1.15, top + 0.9, "DevOps Engineer", "reviews and approves"),
        (1.15, top + 2.9, "SRE / Manager", "monitors outcomes"),
        (11.05, top + 0.9, "External Systems", "Docker, GitHub, Jira"),
        (11.05, top + 2.9, "Scheduler", "triggers checks"),
    ]
    for x, y, name, detail in actors:
        node(slide, Inches(x - 0.85), Inches(y), Inches(1.7), Inches(0.86),
             fill=OVERLAY, shape=MSO_SHAPE.OVAL,
             lines=[(name, 10.5, True, TEXT)])
        box = slide.shapes.add_textbox(
            Inches(x - 1.0), Inches(y + 0.88), Inches(2.0), Inches(0.3))
        label(box, [detail], size=8.5, colour=MUTED, bold=False, align=PP_ALIGN.CENTER)

        if x < 6:
            arrow(slide, Inches(x + 0.85), Inches(y + 0.43), Inches(3.5), Inches(y + 0.43),
                  colour=ACCENT_2, width=1.25)
        else:
            arrow(slide, Inches(x - 0.85), Inches(y + 0.43), Inches(9.85), Inches(y + 0.43),
                  colour=ACCENT_2, width=1.25)
